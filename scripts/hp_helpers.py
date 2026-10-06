from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
from xml.sax.saxutils import escape
import copy

BLUE = "0096D6"; PURPLE = "6B3A97"; DARK = "1A1A1A"; GREY = "535455"; MID = "87898B"
LIGHT = "E5E8E8"; TINT = "E6F4FB"; PALE = "F3F4F4"; WHITE = "FFFFFF"
GREEN = "2E7D32"; AMBER = "C77700"; RED = "C62828"

I = Inches


def rgb(h):
    return RGBColor.from_string(h)


def remove(shape):
    el = shape._element
    el.getparent().remove(el)


def _style_run(r, size=None, bold=None, color=None, italic=None, font=None, link=None):
    if font: r.font.name = font
    if link:
        try:
            r.hyperlink.address = link
        except Exception:
            pass
    if size: r.font.size = Pt(size)
    if bold is not None: r.font.bold = bold
    if italic is not None: r.font.italic = italic
    if color: r.font.color.rgb = rgb(color)


def set_paras(tf, paras, size=11, color=DARK, align=PP_ALIGN.LEFT, space_after=0, line=None):
    """paras: list of items; each item is str, or dict(runs=[(text,{...})], bullet=bool, align, space_after, size)"""
    tf.clear()
    first = True
    for item in paras:
        if isinstance(item, str):
            item = {"runs": [(item, {})]}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = item.get("align", align)
        sa = item.get("space_after", space_after)
        if sa: p.space_after = Pt(sa)
        p.space_before = Pt(item.get("space_before", 0))
        if line: p.line_spacing = line
        if item.get("bullet"):
            pPr = p._p.get_or_add_pPr()
            lvl = item.get("level", 0)
            pPr.set("marL", str(int(Inches(0.2 + 0.25 * lvl))))
            pPr.set("indent", str(-int(Inches(0.2))))
            for tag in ("a:buNone", "a:buChar", "a:buFont"):
                for e in pPr.findall(qn(tag)): pPr.remove(e)
            bf = parse_xml('<a:buFont xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="Arial"/>')
            ch = "&#8226;" if item.get("level", 0) == 0 else "&#8211;"
            bc = parse_xml(f'<a:buChar xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" char="{ch}"/>')
            # insert before defRPr / tabLst if present
            anchor = pPr.find(qn("a:tabLst"))
            if anchor is None: anchor = pPr.find(qn("a:defRPr"))
            if anchor is not None:
                anchor.addprevious(bf); anchor.addprevious(bc)
            else:
                pPr.append(bf); pPr.append(bc)
        else:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", "0"); pPr.set("indent", "0")
            if pPr.find(qn("a:buNone")) is None:
                bn = parse_xml('<a:buNone xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>')
                anchor = pPr.find(qn("a:defRPr"))
                if anchor is not None: anchor.addprevious(bn)
                else: pPr.append(bn)
        for text, fmt in item["runs"]:
            r = p.add_run()
            r.text = text
            _style_run(r, fmt.get("size", item.get("size", size)), fmt.get("bold", item.get("bold", False)),
                       fmt.get("color", item.get("color", color)), fmt.get("italic"), fmt.get("font"), fmt.get("link"))
    return tf


def text(slide, x, y, w, h, paras, size=11, color=DARK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         margin=0, fill=None, space_after=0, line=None):
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", I(margin))
    tf.vertical_anchor = anchor
    if fill:
        tb.fill.solid(); tb.fill.fore_color.rgb = rgb(fill)
    set_paras(tf, paras if isinstance(paras, list) else [paras], size, color, align, space_after, line)
    return tb


def _noeffect(shape):
    for er in shape._element.iter(qn("a:effectRef")):
        er.set("idx", "0")


def box(slide, x, y, w, h, paras=None, fill=None, line=None, line_w=1.5, dash=False, shape="rect",
        size=10, color=DARK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.05, radius=0.12, bold=False):
    st = {"rect": MSO_SHAPE.RECTANGLE, "rrect": MSO_SHAPE.ROUNDED_RECTANGLE, "oval": MSO_SHAPE.OVAL}[shape]
    s = slide.shapes.add_shape(st, I(x), I(y), I(w), I(h))
    if shape == "rrect":
        s.adjustments[0] = radius
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line); s.line.width = Pt(line_w)
        if dash: s.line.dash_style = MSO_LINE.DASH
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    _noeffect(s)
    tf = s.text_frame
    tf.word_wrap = True
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", I(margin))
    tf.vertical_anchor = anchor
    if paras is not None:
        if isinstance(paras, str):
            paras = [{"runs": [(paras, {"bold": bold})]}]
        set_paras(tf, paras, size, color, align)
    return s


def num(slide, cx, cy, n, d=0.24, fill=GREY, size=9):
    return box(slide, cx - d / 2, cy - d / 2, d, d, [{"runs": [(str(n), {"bold": True})]}], fill=fill,
               shape="oval", size=size, color=WHITE, margin=0)


def line(slide, x1, y1, x2, y2, color=DARK, w=1.25, dash=False, arrow=True, head=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    _noeffect(c)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    if dash:
        c.line.dash_style = MSO_LINE.DASH
    ln = c.line._get_or_add_ln()
    if head:
        ln.append(parse_xml('<a:headEnd xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" type="triangle" w="med" len="med"/>'))
    if arrow:
        ln.append(parse_xml('<a:tailEnd xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" type="triangle" w="med" len="med"/>'))
    return c


def set_title(slide, section):
    t = slide.shapes.title
    tf = t.text_frame
    txBody = tf._txBody
    for p in txBody.findall(qn("a:p")):
        txBody.remove(p)
    x = ('<a:p xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
         '<a:r><a:rPr lang="en-US" dirty="0"/><a:t>Systems &amp; Infrastructure | </a:t></a:r>'
         '<a:r><a:rPr lang="en-US" dirty="0"><a:solidFill><a:schemeClr val="accent1"/></a:solidFill></a:rPr>'
         f'<a:t>{escape(section)}</a:t></a:r></a:p>')
    txBody.append(parse_xml(x))


def set_slide_num(slide, n):
    for f in slide._element.iter(qn("a:fld")):
        if f.get("type") == "slidenum":
            f.find(qn("a:t")).text = str(n)


A_NS = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
BORDER = ('<a:lnB w="12700" cap="flat" cmpd="sng" algn="ctr"><a:solidFill><a:srgbClr val="D9D9D9"/></a:solidFill>'
          '<a:prstDash val="solid"/><a:round/><a:headEnd type="none" w="med" len="med"/><a:tailEnd type="none" w="med" len="med"/></a:lnB>')
NOB = lambda side: f'<a:{side} w="0"><a:noFill/></a:{side}>'


def _cell(txt, size, color, bold, fill=None, align="l"):
    paras = txt if isinstance(txt, list) else [txt]
    ps = ""
    for ptxt in paras:
        runs = ptxt if isinstance(ptxt, list) else [(ptxt, {})]
        rs = ""
        for t, f in runs:
            b = "1" if f.get("bold", bold) else "0"
            c = f.get("color", color)
            rs += (f'<a:r><a:rPr lang="en-US" sz="{int(f.get("size", size) * 100)}" b="{b}" dirty="0">'
                   f'<a:solidFill><a:srgbClr val="{c}"/></a:solidFill></a:rPr><a:t>{escape(t)}</a:t></a:r>')
        ps += f'<a:p><a:pPr algn="{align}"><a:lnSpc><a:spcPct val="95000"/></a:lnSpc></a:pPr>{rs}</a:p>'
    fill_xml = f'<a:solidFill><a:srgbClr val="{fill}"/></a:solidFill>' if fill else "<a:noFill/>"
    return (f'<a:tc><a:txBody><a:bodyPr/><a:lstStyle/>{ps}</a:txBody>'
            f'<a:tcPr marL="73152" marR="73152" marT="54864" marB="54864" anchor="t">'
            f'{NOB("lnL")}{NOB("lnR")}{NOB("lnT")}{BORDER}{fill_xml}</a:tcPr></a:tc>')


def rebuild_table(frame, header, rows, widths, size=10.5, hsize=11, row_h=0.4, x=None, y=None,
                  col_fmt=None, header_h=0.36):
    """Rebuild a table in the template's clean style (accent-blue header text, grey rules, no fill)."""
    col_fmt = col_fmt or {}
    tbl = frame._element.find(".//" + qn("a:tbl"))
    grid = "".join(f'<a:gridCol w="{int(Inches(w))}"/>' for w in widths)
    trs = f'<a:tr h="{int(Inches(header_h))}">' + "".join(_cell(h_, hsize, BLUE, False, align=col_fmt.get(i_, {}).get("align", "l")) for i_, h_ in enumerate(header)) + "</a:tr>"
    for r in rows:
        cells = ""
        for ci, v in enumerate(r):
            fmt = col_fmt.get(ci, {})
            cells += _cell(v, fmt.get("size", size), fmt.get("color", DARK), fmt.get("bold", False),
                           align=fmt.get("align", "l"))
        trs += f'<a:tr h="{int(Inches(row_h))}">{cells}</a:tr>'
    new = parse_xml(f'<a:tbl {A_NS}><a:tblPr firstRow="1" bandRow="0">'
                    '<a:tableStyleId>{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}</a:tableStyleId></a:tblPr>'
                    f'<a:tblGrid>{grid}</a:tblGrid>{trs}</a:tbl>')
    tbl.getparent().replace(tbl, new)
    frame.width = Inches(sum(widths))
    frame.height = Inches(header_h + row_h * len(rows))
    if x is not None: frame.left = Inches(x)
    if y is not None: frame.top = Inches(y)
