#!/usr/bin/env python3
"""md2hp.py - convert a deck-markdown file into an HP-template PowerPoint deck.

Usage:
    python md2hp.py deck.md -o deck.pptx [--template ../assets/hp_template.pptx]

See references/markdown-spec.md for the input format. The script:
  * picks a layout for each slide (or obeys <!-- layout: x -->),
  * auto-sizes text to fit, and splits long bullet lists / tables across slides,
  * prints a WARN line for anything it could not fit, so the caller can shorten
    the markdown and re-run.
"""
import argparse
import copy
import math
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import PP_PLACEHOLDER  # noqa: E402
from hp_helpers import *  # noqa: E402,F401,F403

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_TEMPLATE = os.path.join(HERE, "..", "assets", "hp_template.pptx")

# content area (inches) on a 13.33 x 7.5 slide
X0, Y0, W, Y1 = 0.666, 1.35, 12.0, 6.85
CW = 0.55          # average character width in em (conservative: wider than HP Simplified)
LH = 1.22          # line height factor
WARNINGS = []


def warn(msg):
    WARNINGS.append(msg)


# ----------------------------------------------------------------------------
# text measurement
# ----------------------------------------------------------------------------
def est_lines(text, width_in, pt, bold=False):
    cpl = max(1.0, width_in * 72.0 / (pt * (CW + 0.07 if bold else CW)))
    lines = 0
    for para in str(text).split("\n"):
        words = para.split()
        if not words:
            lines += 1
            continue
        cur = 0
        n = 1
        for w in words:
            L = len(w) + (1 if cur else 0)
            if cur + L > cpl and cur > 0:
                n += 1
                cur = len(w)
            else:
                cur += L
        lines += n
    return lines


def est_h(text, width_in, pt, spacing=1.0, bold=False):
    return est_lines(text, width_in, pt, bold) * pt * LH * spacing / 72.0


def plain(text):
    return re.sub(r"\*\*|__|`|(?<!\w)\*(?!\s)|\[([^\]]+)\]\([^)]+\)", lambda m: m.group(1) or "", text)


# ----------------------------------------------------------------------------
# inline markdown -> runs
# ----------------------------------------------------------------------------
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|(?<![\w*])\*[^*\s][^*]*?\*(?![\w*]))")


def runs(text, base=None):
    base = dict(base or {})
    out = []
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], dict(base)))
        tok = m.group(0)
        f = dict(base)
        if tok.startswith("**"):
            f["bold"] = True
            out.append((tok[2:-2], f))
        elif tok.startswith("`"):
            f["font"] = "Consolas"
            f.setdefault("color", GREY)
            out.append((tok[1:-1], f))
        elif tok.startswith("["):
            mm = re.match(r"\[([^\]]+)\]\(([^)]+)\)", tok)
            f["link"] = mm.group(2)
            f["color"] = BLUE
            out.append((mm.group(1), f))
        else:
            f["italic"] = True
            out.append((tok[1:-1], f))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], dict(base)))
    return out or [("", base)]


def split_title_detail(text):
    """'**Name**: detail' or '**Name** - detail' or 'Name: detail' -> (name, detail)."""
    m = re.match(r"^\*\*(.+?)\*\*\s*[:\u2014\u2013-]?\s*(.*)$", text)
    if m:
        return m.group(1).strip(), m.group(2).strip()
    m = re.match(r"^([^:]{1,40}):\s+(.*)$", text)
    if m:
        return m.group(1).strip(), m.group(2).strip()
    return text.strip(), ""


# ----------------------------------------------------------------------------
# markdown parsing
# ----------------------------------------------------------------------------
def parse_front_matter(src):
    meta = {}
    if src.startswith("---"):
        end = src.find("\n---", 3)
        if end != -1:
            for ln in src[3:end].strip().splitlines():
                if ":" in ln:
                    k, v = ln.split(":", 1)
                    meta[k.strip().lower()] = v.strip().strip('"').strip("'")
            src = src[end + 4:]
    return meta, src


def parse_blocks(lines):
    """Parse body lines of one slide into blocks."""
    blocks = []
    i = 0
    n = len(lines)
    while i < n:
        ln = lines[i]
        s = ln.strip()
        if not s:
            i += 1
            continue
        if s.startswith("```"):
            lang = s[3:].strip().lower()
            body = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                body.append(lines[i])
                i += 1
            i += 1
            blocks.append({"t": "flow" if lang == "flow" else "code", "lines": body, "lang": lang})
            continue
        if s.startswith("### "):
            blocks.append({"t": "h3", "text": s[4:].strip()})
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                blocks.append({"t": "table", "header": rows[0], "rows": rows[1:]})
            continue
        if s.startswith(">"):
            body = []
            while i < n and lines[i].strip().startswith(">"):
                body.append(lines[i].strip()[1:].strip())
                i += 1
            blocks.append({"t": "quote", "text": " ".join(b for b in body if b)})
            continue
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", s)
        if m:
            blocks.append({"t": "image", "alt": m.group(1), "path": m.group(2)})
            i += 1
            continue
        if re.match(r"^\s*([-*+]|\d+[.)])\s+", ln):
            ordered = bool(re.match(r"^\s*\d+[.)]\s+", ln))
            items = []
            while i < n and (re.match(r"^\s*([-*+]|\d+[.)])\s+", lines[i]) or
                             (lines[i].startswith("  ") and lines[i].strip() and items)):
                L = lines[i]
                mm = re.match(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$", L)
                if mm:
                    lvl = min(len(mm.group(1).expandtabs(4)) // 2, 2)
                    items.append({"level": 1 if lvl else 0, "text": mm.group(3).strip()})
                else:
                    items[-1]["text"] += " " + L.strip()
                i += 1
            blocks.append({"t": "ol" if ordered else "ul", "items": items})
            continue
        # paragraph
        body = [s]
        i += 1
        while i < n and lines[i].strip() and not re.match(r"^\s*([-*+]|\d+[.)])\s+|^\s*[|>#!`]", lines[i]):
            body.append(lines[i].strip())
            i += 1
        blocks.append({"t": "p", "text": " ".join(body)})
    return blocks


def parse(src):
    meta, src = parse_front_matter(src)
    slides = []
    cur = None
    for raw in src.splitlines():
        if raw.startswith("# ") and not raw.startswith("## "):
            cur = {"kind": "section", "title": raw[2:].strip(), "lines": []}
            slides.append(cur)
        elif raw.startswith("## "):
            cur = {"kind": "content", "title": raw[3:].strip(), "lines": []}
            slides.append(cur)
        elif cur is not None:
            cur["lines"].append(raw)
    for sl in slides:
        body = "\n".join(sl["lines"])
        sl["layout"] = None
        m = re.search(r"<!--\s*layout:\s*([\w-]+)\s*-->", body)
        if m:
            sl["layout"] = m.group(1).lower()
        notes = re.findall(r"<!--\s*notes:\s*(.*?)-->", body, flags=re.S)
        sl["notes"] = "\n".join(x.strip() for x in notes)
        body = re.sub(r"<!--.*?-->", "", body, flags=re.S)
        sl["blocks"] = parse_blocks(body.splitlines())
        del sl["lines"]
    return meta, slides


# ----------------------------------------------------------------------------
# slide creation
# ----------------------------------------------------------------------------
class Deck:
    def __init__(self, template, meta):
        self.prs = Presentation(template)
        self.meta = meta
        self.prefix = meta.get("title_prefix", "Systems & Infrastructure")
        self.footer = meta.get("footer", "HP Confidential")
        L = {l.name: l for l in self.prs.slide_layouts}
        self.L = L

    def new(self, layout_name, footer=True):
        layout = self.L[layout_name]
        s = self.prs.slides.add_slide(layout)
        if footer:
            for ph in layout.placeholders:
                t = ph.placeholder_format.type
                if t in (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER):
                    el = copy.deepcopy(ph._element)
                    s.shapes._spTree.append(el)
            for sh in s.placeholders:
                if sh.placeholder_format.type == PP_PLACEHOLDER.FOOTER:
                    sh.text_frame.text = self.footer
        return s

    def content(self, title):
        s = self.new("Title Only")
        if self.prefix:
            set_title(s, title, self.prefix)
        else:
            s.shapes.title.text = title
        return s

    def save(self, path):
        for n, sl in enumerate(self.prs.slides, 1):
            set_slide_num(sl, n)
        self.prs.save(path)


def set_title(slide, section, prefix):
    t = slide.shapes.title
    txBody = t.text_frame._txBody
    for p in txBody.findall(qn("a:p")):
        txBody.remove(p)
    x = ('<a:p xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
         f'<a:r><a:rPr lang="en-US" dirty="0"/><a:t>{escape(prefix)} | </a:t></a:r>'
         '<a:r><a:rPr lang="en-US" dirty="0"><a:solidFill><a:schemeClr val="accent1"/></a:solidFill></a:rPr>'
         f'<a:t>{escape(section)}</a:t></a:r></a:p>')
    txBody.append(parse_xml(x))


def add_notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = text


# ----------------------------------------------------------------------------
# common pieces
# ----------------------------------------------------------------------------
def lead_and_callout(slide, blocks, y0=Y0, y1=Y1):
    """Render first paragraph as lead line and blockquote as bottom callout; return remaining area."""
    lead = next((b for b in blocks if b["t"] == "p"), None)
    quote = next((b for b in blocks if b["t"] == "quote"), None)
    if lead:
        h = est_h(plain(lead["text"]), W, 13) + 0.05
        text(slide, X0, y0 - 0.05, W, h, [{"runs": runs(lead["text"])}], size=13, color=GREY)
        y0 += h + 0.2
    if quote:
        qh = max(0.62, est_h(plain(quote["text"]), W - 0.4, 12) + 0.3)
        box(slide, X0, y1 - qh, W, qh, [{"runs": runs(quote["text"], {"color": DARK})}], fill=TINT,
            shape="rrect", radius=0.12, size=12, align=PP_ALIGN.LEFT, margin=0.18)
        y1 -= qh + 0.25
    return y0, y1


def fit_size(height_fn, avail, sizes):
    for sz in sizes:
        if height_fn(sz) <= avail:
            return sz, True
    return sizes[-1], False


def bullets_height(items, width, sz, gap_pt=6):
    h = 0
    for it in items:
        isz = sz if it["level"] == 0 else sz - 1
        h += est_h(plain(it["text"]), width - 0.2 - 0.25 * it["level"], isz) + gap_pt / 72.0
    return h


def bullet_paras(items, sz, gap=6):
    out = []
    for it in items:
        isz = sz if it["level"] == 0 else sz - 1
        out.append({"runs": runs(it["text"]), "bullet": True, "level": it["level"], "size": isz,
                    "space_after": gap})
    return out


def para_items(items, sz, gap=6):
    return [p if not it.get("para") else dict(p, bullet=False) for p, it in zip(bullet_paras(items, sz, gap), items)]


def other_blocks(blocks):
    """Everything except the lead paragraph and the quote."""
    first_p = next((b for b in blocks if b["t"] == "p"), None)
    return [b for b in blocks if b is not first_p and b["t"] != "quote"]


def to_items(blocks):
    items = []
    for b in blocks:
        if b["t"] in ("ul", "ol"):
            items += b["items"]
        elif b["t"] == "p":
            items.append({"level": 0, "text": b["text"], "para": True})
    return items


def groups_of(blocks):
    groups = []
    intro = []
    for b in blocks:
        if b["t"] == "h3":
            groups.append({"heading": b["text"], "blocks": []})
        elif groups:
            groups[-1]["blocks"].append(b)
        else:
            intro.append(b)
    return intro, groups


# ----------------------------------------------------------------------------
# layouts
# ----------------------------------------------------------------------------
def r_title(deck, meta):
    s = deck.new("Title Slide", footer=False)
    s.shapes.title.text = meta.get("title", "Untitled")
    sub = [meta.get("subtitle", "")]
    by = " | ".join(x for x in [meta.get("author") and f"Author: {meta['author']}",
                                meta.get("date") and f"Published: {meta['date']}"] if x)
    sub = [x for x in sub + [by] if x]
    for ph in s.placeholders:
        if ph.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE:
            ph.text_frame.text = "\n".join(sub)


def r_section(deck, sl):
    s = deck.new("Blue Section Header")
    s.shapes.title.text = sl["title"]
    p = next((b for b in sl["blocks"] if b["t"] == "p"), None)
    for ph in s.placeholders:
        if ph.placeholder_format.type == PP_PLACEHOLDER.BODY:
            if p:
                ph.text_frame.text = plain(p["text"])
            else:
                ph._element.getparent().remove(ph._element)
    return s


def r_closing(deck):
    s = deck.new("Closing Option 2")
    for ph in list(s.placeholders):
        if ph.placeholder_format.type == PP_PLACEHOLDER.TITLE:
            ph._element.getparent().remove(ph._element)


def bullet_ctx(sl):
    """(items, ctx_blocks): ctx_blocks hold the lead paragraph and/or quote shown around the bullets."""
    blocks = sl["blocks"]
    prose = all(b["t"] in ("p", "quote") for b in blocks)
    if prose:
        return to_items([b for b in blocks if b["t"] != "quote"]), [b for b in blocks if b["t"] == "quote"]
    return to_items(other_blocks(blocks)), [b for b in blocks if b["t"] in ("p", "quote")][:1] + \
        [b for b in blocks if b["t"] == "quote"][:1]


def r_bullets(deck, sl, items=None, title=None, ctx=None):
    all_items, all_ctx = bullet_ctx(sl)
    items = all_items if items is None else items
    ctx = all_ctx if ctx is None else ctx
    s = deck.content(title or sl["title"])
    y0, y1 = lead_and_callout(s, ctx)
    avail = y1 - y0
    two_col = len(items) >= 8 and sum(len(plain(i["text"])) for i in items) / max(1, len(items)) < 70
    if two_col:
        half = math.ceil(len(items) / 2)
        cols = [items[:half], items[half:]]
        cw = (W - 0.5) / 2
        sz, ok = fit_size(lambda z: max(bullets_height(c, cw, z) for c in cols), avail, [16, 15, 14, 13, 12])
        for k, c in enumerate(cols):
            tb = text(s, X0 + k * (cw + 0.5), y0, cw, avail, para_items(c, sz), size=sz)
    else:
        sz, ok = fit_size(lambda z: bullets_height(items, W, z, 10), avail, [18, 16, 15, 14, 13])
        text(s, X0, y0, W, avail, para_items(items, sz, 10), size=sz)
    if not ok:
        return s, False
    return s, True


def r_table(deck, sl, tbl=None, title=None, intro_blocks=None):
    blocks = sl["blocks"] if intro_blocks is None else intro_blocks
    tbl = tbl or next(b for b in sl["blocks"] if b["t"] == "table")
    s = deck.content(title or sl["title"])
    y0, y1 = lead_and_callout(s, blocks)
    ncol = len(tbl["header"])
    rows = [r + [""] * (ncol - len(r)) for r in tbl["rows"]]
    # column widths proportional to content length (bounded)
    lens = []
    for c in range(ncol):
        L = [len(plain(tbl["header"][c]))] + [len(plain(r[c])) for r in rows]
        lens.append(max(4, min(60, sum(L) / len(L) * 0.6 + max(L) * 0.4)))
    tot = sum(lens)
    widths = [max(0.6, W * l / tot) for l in lens]
    k = W / sum(widths)
    widths = [w * k for w in widths]

    def height(sz):
        h = 0.36
        for r in rows:
            h += max(est_h(plain(r[c]), widths[c] - 0.16, sz) for c in range(ncol)) + 0.14
        return h

    sz, ok = fit_size(height, y1 - y0, [14, 13, 12, 11.5, 11, 10.5, 10, 9.5])
    frame = s.shapes.add_table(len(rows) + 1, ncol, I(X0), I(y0), I(W), I(1))
    fmt = {}
    # bold first column when it looks like a label column
    if ncol >= 3 and all(len(plain(r[0])) < 45 for r in rows):
        fmt[0] = {"bold": True}
    for c in range(ncol):
        vals = [plain(r[c]).strip().lower() for r in rows]
        if vals and all(len(v) <= 6 for v in vals):
            fmt[c] = dict(fmt.get(c, {}), align="ctr")
    cell = lambda v: [[(t, f) for t, f in runs(v)]]
    rebuild_table(frame, [plain(h).upper() for h in tbl["header"]], [[cell(v) for v in r] for r in rows],
                  widths, size=sz, hsize=min(11.5, sz), row_h=0.3, header_h=0.36, x=X0, y=y0, col_fmt=fmt)
    return s, ok


def r_columns(deck, sl):
    intro, groups = groups_of(other_blocks(sl["blocks"]))
    s = deck.content(sl["title"])
    y0, y1 = lead_and_callout(s, sl["blocks"])
    n = len(groups)
    if n == 2:
        gap = 0.6
        cw = (W - gap) / 2
        line(s, X0 + cw + gap / 2, y0, X0 + cw + gap / 2, y1, color="B9B8BB", w=1, dash=True, arrow=False)
        pad = 0
        card = False
    else:
        gap = 0.25
        cw = (W - gap * (n - 1)) / n
        pad = 0.2
        card = True
    avail = y1 - y0 - 0.45 - 2 * pad
    items_by = [to_items(g["blocks"]) for g in groups]
    sz, ok = fit_size(lambda z: max(bullets_height(it, cw - 2 * pad, z) for it in items_by) if items_by else 0,
                      avail, [16, 15, 14, 13, 12, 11, 10.5])
    ch = max(bullets_height(it, cw - 2 * pad, sz) for it in items_by) + 0.48 + 2 * pad + 0.15
    ch = min(y1 - y0, max(2.4, ch))
    for k, g in enumerate(groups):
        x = X0 + k * (cw + gap)
        if card:
            box(s, x, y0, cw, ch, fill=PALE, shape="rrect", radius=0.06)
        text(s, x + pad, y0 + pad, cw - 2 * pad, 0.4, [{"runs": runs(g["heading"], {"bold": True})}],
             size=15, color=BLUE)
        items = items_by[k]
        paras = [p if not it.get("para") else dict(p, bullet=False) for p, it in zip(bullet_paras(items, sz), items)]
        text(s, x + pad, y0 + pad + 0.48, cw - 2 * pad, avail, paras, size=sz)
    return s, ok


def r_process(deck, sl):
    blocks = sl["blocks"]
    ol = next(b for b in blocks if b["t"] == "ol")
    steps = [split_title_detail(i["text"]) for i in ol["items"] if i["level"] == 0]
    n = len(steps)
    if n > 6:
        return r_steps_vertical(deck, sl, steps)
    s = deck.content(sl["title"])
    y0, y1 = lead_and_callout(s, blocks)
    colw = W / n
    d = 0.9
    dw = colw - 0.3
    sz, ok = fit_size(lambda z: max(est_h(t[1], dw, z) for t in steps) + 0.5 + d + 0.4, y1 - y0,
                      [15, 14, 13, 12, 11, 10.5, 10])
    blk = d + 0.2 + max(est_h(t[0], dw, 15, bold=True) for t in steps) + max(est_h(t[1], dw, sz) for t in steps) + 0.2
    top = y0 + max(0.1, (y1 - y0 - blk) / 2 - 0.2)
    cy = top + d / 2
    for k, (name, det) in enumerate(steps):
        cx = X0 + colw * k + colw / 2
        box(s, cx - d / 2, cy - d / 2, d, d, [{"runs": [(str(k + 1), {"bold": True})]}], fill=BLUE,
            shape="oval", size=22, color=WHITE, margin=0)
        if k < n - 1:
            line(s, cx + d / 2 + 0.12, cy, cx + colw - d / 2 - 0.12, cy, color=MID, w=1.5)
        nh = max(est_h(name, dw, 15, bold=True) for name, _ in steps) + 0.08
        text(s, cx - dw / 2, cy + d / 2 + 0.15, dw, nh, [{"runs": runs(name, {"bold": True})}], size=15,
             align=PP_ALIGN.CENTER)
        if det:
            text(s, cx - dw / 2, cy + d / 2 + 0.2 + nh, dw, y1 - (cy + d / 2 + 0.2 + nh), [{"runs": runs(det)}],
                 size=sz, color=GREY, align=PP_ALIGN.CENTER)
    return s, ok


def r_steps_vertical(deck, sl, steps):
    s = deck.content(sl["title"])
    y0, y1 = lead_and_callout(s, sl["blocks"])
    n = len(steps)
    cols = 1 if n <= 7 else 2
    per = math.ceil(n / cols)
    cw = (W - 0.3 * (cols - 1)) / cols
    pitch = min(0.85, (y1 - y0) / per)
    rh = pitch - 0.1
    ok = rh >= 0.45
    for k, (name, det) in enumerate(steps):
        c, r = divmod(k, per)
        x, y = X0 + c * (cw + 0.3), y0 + r * pitch
        box(s, x, y, cw, rh, fill=PALE, shape="rrect", radius=0.12)
        num(s, x + 0.32, y + rh / 2, k + 1, d=0.32, fill=BLUE, size=11)
        paras = [{"runs": runs(name, {"bold": True}), "size": 12}]
        if det:
            paras.append({"runs": runs(det), "size": 10.5, "color": GREY})
        text(s, x + 0.65, y + 0.04, cw - 0.8, rh - 0.08, paras, anchor=MSO_ANCHOR.MIDDLE)
    return s, ok


TAG = re.compile(r"^\[([A-Za-z][A-Za-z \-]{1,20})\]\s*")


def r_compare(deck, sl):
    intro, groups = groups_of(other_blocks(sl["blocks"]))
    g1, g2 = groups[0], groups[1]
    left = [i for b in g1["blocks"] if b["t"] in ("ul", "ol") for i in b["items"] if i["level"] == 0]
    right = [i for b in g2["blocks"] if b["t"] in ("ul", "ol") for i in b["items"] if i["level"] == 0]
    n = max(len(left), len(right))
    s = deck.content(sl["title"])
    y0, y1 = lead_and_callout(s, sl["blocks"])
    tags = []
    rt = []
    for it in right:
        m = TAG.match(it["text"])
        tags.append(m.group(1).upper() if m else None)
        rt.append(it["text"][m.end():] if m else it["text"])
    has_tags = any(tags)
    MW = 1.5 if has_tags else 0.5
    LW = (W - MW) / 2
    LX, MX, RX = X0, X0 + LW, X0 + LW + MW
    text(s, LX, y0, LW, 0.3, [{"runs": runs(g1["heading"].upper(), {"bold": True})}], size=11, color=GREY)
    text(s, RX, y0, LW, 0.3, [{"runs": runs(g2["heading"].upper(), {"bold": True})}], size=11, color=BLUE)
    if has_tags:
        text(s, MX, y0, MW, 0.3, [{"runs": [("CHANGE", {"bold": True})]}], size=11, color=GREY, align=PP_ALIGN.CENTER)
    top = y0 + 0.4
    pitch = (y1 - top) / max(n, 1)
    pitch = min(pitch, 0.95)
    rh = pitch - 0.1
    ok = True
    for k in range(n):
        y = top + k * pitch
        for side, (x, items, fill, nf) in enumerate([(LX, left, PALE, GREY), (RX, rt, TINT, BLUE)]):
            src = items[k] if k < len(items) else None
            txt = (src["text"] if isinstance(src, dict) else src) if src else None
            if txt and txt.strip() not in ("-", "\u2014", "none", "None"):
                box(s, x, y, LW, rh, fill=fill, shape="rrect", radius=0.12)
                num(s, x + 0.3, y + rh / 2, k + 1, d=0.3, fill=nf, size=10)
                name, det = split_title_detail(txt)
                paras = [{"runs": runs(name, {"bold": True}), "size": 11}]
                if det:
                    paras.append({"runs": runs(det), "size": 9.5, "color": GREY})
                need = est_h(name, LW - 0.75, 11) + (est_h(det, LW - 0.75, 9.5) if det else 0)
                ok = ok and need <= rh
                text(s, x + 0.6, y + 0.04, LW - 0.75, rh - 0.08, paras, anchor=MSO_ANCHOR.MIDDLE)
            else:
                box(s, x, y, LW, rh, line="B9B8BB", line_w=1, dash=True, shape="rrect", radius=0.12)
                text(s, x + 0.6, y, LW - 0.75, rh, [{"runs": [("\u2014  no equivalent step", {"italic": True})]}],
                     size=10, color=MID, anchor=MSO_ANCHOR.MIDDLE)
        if has_tags and tags[k] if k < len(tags) else False:
            t = tags[k]
            f, c = (LIGHT, GREY) if t in ("UNCHANGED", "SAME") else ((("99D5EF", DARK)) if t in ("ENHANCED", "CHANGED") else (BLUE, WHITE))
            box(s, MX + 0.08, y + rh / 2 - 0.16, MW - 0.16, 0.32, [{"runs": [(t, {"bold": True})]}], fill=f,
                shape="rrect", radius=0.5, size=8.5, color=c, margin=0.02)
        elif not has_tags:
            line(s, MX + 0.1, y + rh / 2, MX + MW - 0.1, y + rh / 2, color=MID, w=1.5)
    return s, ok


def r_stats(deck, sl):
    items = [i for i in to_items(other_blocks(sl["blocks"])) if i["level"] == 0]
    s = deck.content(sl["title"])
    y0, y1 = lead_and_callout(s, sl["blocks"])
    n = len(items)
    per = min(4, n)
    rows = math.ceil(n / per)
    cw = (W - 0.25 * (per - 1)) / per
    rh = min(2.2, (y1 - y0 - 0.25 * (rows - 1)) / rows)
    blk = rows * rh + 0.25 * (rows - 1)
    top = y0 + max(0, (y1 - y0 - blk) / 2 - 0.3)
    ok = True
    for k, it in enumerate(items):
        r, c = divmod(k, per)
        x, y = X0 + c * (cw + 0.25), top + r * (rh + 0.25)
        m = re.match(r"^\*\*(.+?)\*\*\s*[:\u2014\u2013-]?\s*(.*)$", it["text"])
        val, lab = (m.group(1), m.group(2)) if m else (it["text"], "")
        box(s, x, y, cw, rh, fill=PALE, shape="rrect", radius=0.08)
        vs = 40 if len(val) <= 6 else 28 if len(val) <= 10 else 20
        text(s, x + 0.2, y + 0.2, cw - 0.4, 0.8, [{"runs": [(val, {"bold": True})]}], size=vs, color=BLUE)
        text(s, x + 0.2, y + 1.05, cw - 0.4, rh - 1.15, [{"runs": runs(lab)}], size=13, color=GREY)
        ok = ok and est_h(plain(lab), cw - 0.4, 13) <= rh - 1.15
    return s, ok


def r_quote(deck, sl):
    q = next(b for b in sl["blocks"] if b["t"] == "quote")
    s = deck.content(sl["title"])
    txt = q["text"]
    m = re.match(r"^(.*?)\s*[\u2014-]{1,2}\s*([^\u2014-]{2,60})$", txt)
    body, who = (m.group(1), m.group(2)) if m and len(m.group(1)) > 20 else (txt, "")
    sz, ok = fit_size(lambda z: est_h(plain(body), 10.5, z), 3.6, [28, 24, 22, 20, 18])
    text(s, X0 + 0.75, 2.0, 10.5, 3.6, [{"runs": runs("\u201c" + body + "\u201d")}], size=sz, color=DARK,
         anchor=MSO_ANCHOR.MIDDLE)
    if who:
        text(s, X0 + 0.75, 5.75, 10.5, 0.4, [{"runs": [("\u2014 " + who, {"bold": True})]}], size=14, color=BLUE)
    line(s, X0 + 0.4, 2.1, X0 + 0.4, 5.5, color=BLUE, w=4, arrow=False)
    return s, ok


def r_image(deck, sl, base_dir):
    im = next(b for b in sl["blocks"] if b["t"] == "image")
    s = deck.content(sl["title"])
    y0, y1 = lead_and_callout(s, sl["blocks"])
    path = im["path"] if os.path.isabs(im["path"]) else os.path.join(base_dir, im["path"])
    if not os.path.exists(path):
        warn(f"image not found: {path}")
        box(s, X0, y0, W, y1 - y0, f"Image not found: {im['path']}", line=MID, dash=True, size=12, color=MID)
        return s, True
    from PIL import Image as PImage
    w, h = PImage.open(path).size
    cap = 0.35 if im["alt"] else 0
    aw, ah = W, y1 - y0 - cap
    sc = min(aw / w, ah / h)
    iw, ih = w * sc, h * sc
    s.shapes.add_picture(path, I(X0 + (aw - iw) / 2), I(y0), I(iw), I(ih))
    if im["alt"]:
        text(s, X0, y0 + ih + 0.08, W, 0.3, im["alt"], size=10, color=GREY, align=PP_ALIGN.CENTER)
    return s, True


def parse_flow(lines):
    nodes, edges = {}, []

    def node(tok):
        tok = tok.strip()
        tags = set(t.strip().lower() for t in re.findall(r"\[([^\]]+)\]", tok))
        name = re.sub(r"\s*\[[^\]]+\]", "", tok).strip()
        n = nodes.setdefault(name, {"name": name, "it": False, "new": False})
        n["it"] |= "it" in tags
        n["new"] |= "new" in tags
        return name

    for ln in lines:
        ln = ln.strip()
        if not ln or ln.startswith("#"):
            continue
        label = ""
        if ":" in ln and re.search(r"-->|->", ln.split(":", 1)[0]):
            ln, label = ln.split(":", 1)
            label = label.strip()
        parts = re.split(r"(-->|->)", ln)
        if len(parts) == 1:
            node(parts[0])
            continue
        prev = node(parts[0])
        for i in range(1, len(parts), 2):
            nxt = node(parts[i + 1])
            edges.append({"a": prev, "b": nxt, "new": parts[i] == "-->", "label": label if i == len(parts) - 2 else ""})
            prev = nxt
    return nodes, edges


def r_flow(deck, sl):
    fb = next(b for b in sl["blocks"] if b["t"] == "flow")
    nodes, edges = parse_flow(fb["lines"])
    s = deck.content(sl["title"])
    y0, y1 = lead_and_callout(s, sl["blocks"])
    # layer by longest path (ignore back edges)
    order = list(nodes)
    idx = {n: i for i, n in enumerate(order)}
    level = {n: 0 for n in order}
    for _ in range(len(order)):
        changed = False
        for e in edges:
            # forward edges only (b first appears after a); back edges don't push layers
            if idx[e["b"]] > idx[e["a"]] and level[e["b"]] < level[e["a"]] + 1:
                level[e["b"]] = level[e["a"]] + 1
                changed = True
        if not changed:
            break
    ncol = max(level.values()) + 1
    cols = [[n for n in order if level[n] == c] for c in range(ncol)]
    nrow = max(len(c) for c in cols)
    legend = any(n["new"] or n["it"] for n in nodes.values()) or any(e["new"] for e in edges)
    ya = y1 - (0.4 if legend else 0)
    has_lab = any(e["label"] for e in edges)
    mingap = 1.5 if has_lab else 0.7
    bw = max(1.25, min(2.4, (W - mingap * (ncol - 1)) / ncol))
    bh = 0.8
    gapx = (W - bw * ncol) / max(1, ncol - 1) if ncol > 1 else 0
    pos = {}
    for c, col in enumerate(cols):
        x = X0 + c * (bw + gapx) if ncol > 1 else X0 + (W - bw) / 2
        pitch = min(2.0, (ya - y0) / nrow)
        span = pitch * len(col)
        top = y0 + ((ya - y0) - span) / 2
        for r, n in enumerate(col):
            y = top + r * pitch + (pitch - bh) / 2
            pos[n] = (x, y)
    pair_count, pair_seen = {}, {}
    for e in edges:
        key = tuple(sorted((e["a"], e["b"])))
        pair_count[key] = pair_count.get(key, 0) + 1
    for e in edges:
        key = tuple(sorted((e["a"], e["b"])))
        k = pair_seen.get(key, 0)
        pair_seen[key] = k + 1
        off = (k - (pair_count[key] - 1) / 2) * 0.2
        (ax, ay), (bx, by) = pos[e["a"]], pos[e["b"]]
        ay += off
        by += off
        col = BLUE if e["new"] else DARK
        if abs(ax - bx) < 0.01:
            x1, y1_, x2, y2 = ax + bw / 2, ay + (bh if by > ay else 0), bx + bw / 2, by + (0 if by > ay else bh)
        elif bx > ax:
            x1, y1_, x2, y2 = ax + bw, ay + bh / 2, bx, by + bh / 2
        else:
            x1, y1_, x2, y2 = ax, ay + bh / 2, bx + bw, by + bh / 2
        line(s, x1, y1_, x2, y2, color=col, dash=e["new"], w=1.5)
        if e["label"]:
            mx, my = (x1 + x2) / 2, (y1_ + y2) / 2
            lw = max(1.2, min(abs(x2 - x1) - 0.1, 2.4))
            ly = my - 0.27 if off <= 0 else my + 0.04
            text(s, mx - lw / 2, ly, lw, 0.24, e["label"], size=9, color=col, align=PP_ALIGN.CENTER, fill=WHITE)
    for n, (x, y) in pos.items():
        nd = nodes[n]
        box(s, x, y, bw, bh, [{"runs": runs(n)}], fill=WHITE, line=PURPLE if nd["it"] else BLUE, dash=nd["new"],
            size=12 if len(n) < 18 else 11, color=DARK, margin=0.05)
    if legend:
        lx, ly = X0, y1 - 0.22
        for kind, colr, dsh, lab in [("box", BLUE, False, "S&I system"), ("box", PURPLE, False, "IT system"),
                                     ("box", BLUE, True, "New component"), ("ln", DARK, False, "Current flow"),
                                     ("ln", BLUE, True, "New flow")]:
            if kind == "box":
                box(s, lx, ly, 0.28, 0.18, line=colr, dash=dsh, line_w=1.25)
            else:
                line(s, lx, ly + 0.09, lx + 0.28, ly + 0.09, color=colr, dash=dsh, w=1.25)
            text(s, lx + 0.34, ly - 0.02, 1.3, 0.22, lab, size=9, color=GREY)
            lx += 1.6
    ok = nrow * bh <= (ya - y0)
    return s, ok


def r_code(deck, sl):
    cb = next(b for b in sl["blocks"] if b["t"] == "code")
    s = deck.content(sl["title"])
    y0, y1 = lead_and_callout(s, sl["blocks"])
    body = "\n".join(cb["lines"])
    n = len(cb["lines"])
    sz, ok = fit_size(lambda z: n * z * 1.25 / 72 + 0.3, y1 - y0, [14, 12, 11, 10, 9])
    b = box(s, X0, y0, W, min(y1 - y0, n * sz * 1.25 / 72 + 0.35), fill=PALE, shape="rrect", radius=0.03,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.15)
    set_paras(b.text_frame, [{"runs": [(ln or " ", {"font": "Consolas"})]} for ln in cb["lines"]], size=sz, color=DARK)
    return s, ok


# ----------------------------------------------------------------------------
# layout choice + pagination
# ----------------------------------------------------------------------------
def choose(sl):
    if sl["layout"]:
        return sl["layout"]
    b = sl["blocks"]
    types = [x["t"] for x in b]
    intro, groups = groups_of(b)
    if "flow" in types:
        return "flow"
    if "image" in types:
        return "image"
    if "table" in types:
        return "table"
    if len(groups) == 2 and all(any(x["t"] in ("ol", "ul") for x in g["blocks"]) for g in groups) and \
            re.search(r"current|as-is|before|today|old", groups[0]["heading"], re.I) and \
            re.search(r"proposed|to-be|after|target|new|future", groups[1]["heading"], re.I):
        return "compare"
    if 2 <= len(groups) <= 4:
        return "columns"
    if "code" in types:
        return "code"
    lists = [x for x in b if x["t"] in ("ul", "ol")]
    if types and set(types) <= {"quote"}:
        return "quote"
    if len(lists) == 1 and lists[0]["t"] == "ol" and 3 <= len([i for i in lists[0]["items"] if i["level"] == 0]) <= 6:
        return "process"
    if lists and all(re.match(r"^\*\*[^*]{0,10}\d[^*]{0,10}\*\*", i["text"]) for x in lists for i in x["items"]) and \
            len(lists[0]["items"]) <= 8 and all(len(plain(i["text"])) < 90 for i in lists[0]["items"]):
        return "stats"
    return "bullets"


def drop_last(deck):
    lst = deck.prs.slides._sldIdLst
    sid = lst[-1]
    deck.prs.part.drop_rel(sid.rId)
    lst.remove(sid)


def split_even(items, pages):
    """Split a list into `pages` nearly equal chunks without starting a chunk on a sub-bullet."""
    n = len(items)
    cuts, out, start = [round(n * k / pages) for k in range(1, pages)], [], 0
    for c in cuts:
        while c < n and isinstance(items[c], dict) and items[c].get("level", 0) > 0:
            c += 1
        if c > start:
            out.append(items[start:c])
            start = c
    out.append(items[start:])
    return [o for o in out if o]


def render_slide(deck, sl, base_dir):
    lay = choose(sl)
    title = sl["title"]
    try:
        if lay == "bullets":
            items, ctx = bullet_ctx(sl)
            s, ok = r_bullets(deck, sl)
            if not ok and len(items) > 3:
                drop_last(deck)
                lead = [c for c in ctx if c["t"] == "p"]
                quote = [c for c in ctx if c["t"] == "quote"]
                for pages in range(2, 9):
                    chunks = split_even(items, pages)
                    made, allok = 0, True
                    for k, ch in enumerate(chunks, 1):
                        cx = (lead if k == 1 else []) + (quote if k == len(chunks) else [])
                        s, ok = r_bullets(deck, sl, ch, f"{title} ({k}/{len(chunks)})", cx)
                        made += 1
                        allok = allok and ok
                        if k < len(chunks):
                            add_notes(s, sl["notes"])
                    if allok or pages == 8:
                        ok = allok
                        break
                    for _ in range(made):
                        drop_last(deck)
        elif lay == "table":
            tbl = next(b for b in sl["blocks"] if b["t"] == "table")
            s, ok = r_table(deck, sl)
            if not ok and len(tbl["rows"]) > 2:
                drop_last(deck)
                rows = tbl["rows"]
                for pages in range(2, 11):
                    chunks = split_even(rows, pages)
                    made, allok = 0, True
                    for k, ch in enumerate(chunks, 1):
                        intro = [b for b in sl["blocks"] if b["t"] == "p"][:1] if k == 1 else []
                        intro += [b for b in sl["blocks"] if b["t"] == "quote"] if k == len(chunks) else []
                        s, ok = r_table(deck, sl, dict(tbl, rows=ch), f"{title} ({k}/{len(chunks)})",
                                        intro_blocks=intro)
                        made += 1
                        allok = allok and ok
                    if allok or pages == 10:
                        ok = allok
                        break
                    for _ in range(made):
                        drop_last(deck)
        elif lay == "columns":
            s, ok = r_columns(deck, sl)
        elif lay == "compare":
            s, ok = r_compare(deck, sl)
        elif lay == "process":
            s, ok = r_process(deck, sl)
        elif lay == "stats":
            s, ok = r_stats(deck, sl)
        elif lay == "quote":
            s, ok = r_quote(deck, sl)
        elif lay == "image":
            s, ok = r_image(deck, sl, base_dir)
        elif lay == "flow":
            s, ok = r_flow(deck, sl)
        elif lay == "code":
            s, ok = r_code(deck, sl)
        else:
            warn(f"slide '{title}': unknown layout '{lay}', used bullets")
            s, ok = r_bullets(deck, sl)
    except (StopIteration, IndexError, KeyError) as e:
        warn(f"slide '{title}': layout '{lay}' did not match content ({e!r}); used bullets")
        s, ok = r_bullets(deck, sl)
        lay = "bullets"
    if not ok:
        warn(f"slide '{title}' ({lay}): content may overflow; shorten text or split the slide")
    add_notes(s, sl["notes"])
    return lay


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("markdown")
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--template", default=DEFAULT_TEMPLATE)
    a = ap.parse_args()
    src = open(a.markdown, encoding="utf-8").read()
    meta, slides = parse(src)
    deck = Deck(a.template, meta)
    r_title(deck, meta)
    plan = []
    for sl in slides:
        if sl["kind"] == "section":
            s = r_section(deck, sl)
            add_notes(s, sl["notes"])
            plan.append(("section", sl["title"]))
        else:
            plan.append((render_slide(deck, sl, os.path.dirname(os.path.abspath(a.markdown))), sl["title"]))
    if meta.get("closing", "true").lower() != "false":
        r_closing(deck)
    deck.save(a.out)
    print(f"Saved {a.out}: {len(deck.prs.slides)} slides")
    for lay, t in plan:
        print(f"  [{lay:8}] {t}")
    for w_ in WARNINGS:
        print("WARN", w_)
    return 1 if WARNINGS else 0


if __name__ == "__main__":
    sys.exit(main())
